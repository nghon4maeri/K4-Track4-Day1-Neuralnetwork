from __future__ import annotations
import json
import os
import openpyxl
from pathlib import Path

def save_result(result: dict, results_dir: str = "../results") -> str:
    os.makedirs(results_dir, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    out_path = f"{results_dir}/{exp_id}.json"
    
    save_dict = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"]
    }
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(save_dict, f, indent=2)
    return out_path

def load_results(results_dir: str = "../results") -> list[dict]:
    results = []
    if not os.path.exists(results_dir):
        return results
    for filename in sorted(os.listdir(results_dir)):
        if filename.endswith(".json"):
            with open(f"{results_dir}/{filename}", "r", encoding="utf-8") as f:
                results.append(json.load(f))
    return results

def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    cfg = result["cfg"]
    summary = result["summary"]
    
    row = {
        "exp_id": cfg["exp_id"],
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg["loss"],
        "optimizer": cfg["optimizer"],
        "lr": cfg["lr"],
        "weight_decay": cfg["weight_decay"],
        "batch": cfg["batch"],
        "epochs": cfg["epochs"],
        "hidden": str(cfg["hidden"]),
        "dropout": cfg["dropout"],
        "clip_norm": cfg["clip_norm"] if cfg["clip_norm"] is not None else "",
        "precision": cfg["precision"],
        "init": cfg["init"],
        "seed": cfg["seed"],
        "step0_loss": summary.get("step0_loss", ""),
        "best_val_loss": summary.get("best_val_loss", ""),
        "best_epoch": summary.get("best_epoch", ""),
        "final_train_loss": summary.get("final_train_loss", ""),
        "final_val_loss": summary.get("final_val_loss", ""),
        "val_acc": summary.get("val_acc", ""),
        "val_macro_f1": summary.get("val_macro_f1", ""),
        "time_per_epoch_s": summary.get("time_per_epoch_s", ""),
        "peak_mem_MB": summary.get("peak_mem_MB", ""),
        "diverged": summary.get("diverged", False),
        "figure_file": f"figures/{cfg['exp_id']}.png",
        "notes": notes
    }
    
    if eval_scores is not None:
        row["eval_acc"] = eval_scores.get("accuracy", "")
        row["eval_macro_f1"] = eval_scores.get("macro_f1", "")
    else:
        row["eval_acc"] = ""
        row["eval_macro_f1"] = ""
        
    return row

def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    
    header = {cell.value: idx for idx, cell in enumerate(ws[1]) if cell.value}
    skip_cols = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
    
    start_row = 2
    for r_idx, row_dict in enumerate(rows):
        excel_row = start_row + r_idx
        for key, val in row_dict.items():
            if key in header and key not in skip_cols:
                col_idx = header[key] + 1
                ws.cell(row=excel_row, column=col_idx, value=val)
                
    wb.save(out_path)
